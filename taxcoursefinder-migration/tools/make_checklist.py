import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

S = sys.argv[1]
P = sorted(json.load(open(S + '/pf/providers_now.json'))['providers'], key=lambda p: p['name'].lower())
F = 'Arial'
NAVY, OLIVE, GREY = '0B2E40', '95AB3B', '5A6E78'
thin = Side(style='thin', color='D7DFE2')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# (group, header, kind, definition, current-site key)
COLS = [
    ('Offers & Pricing', 'Offers 20-Hour CE', 'yn', 'Sells a CTEC-approved 20-hour continuing education (renewal) course.', ('flag', 'offers_20_hour_ce')),
    ('Offers & Pricing', 'Offers 60-Hour QE', 'yn', 'Sells a CTEC-approved 60-hour qualifying education course for new preparers.', ('flag', 'offers_60_hour_qe')),
    ('Offers & Pricing', '20-Hour course bundle price ($)', 'money', 'Price of the complete 20-hour package as listed on the site today. Leave blank if not published.', ('price', 'ce')),
    ('Offers & Pricing', '60-Hour course bundle price ($)', 'money', 'Price of the complete 60-hour package as listed on the site today. Leave blank if not published.', ('price', 'qe')),
    ('Offers & Pricing', 'Coupon or sale (code / details)', 'text', 'Any public coupon code or sale, e.g. "SAVE20 – 20% off, expires 12/31/2026". Leave blank if none.', ('offers', None)),
    ('Delivery / Format', 'Online Course', 'yn', 'Course can be taken fully online.', ('feat', 'online_course_available')),
    ('Delivery / Format', 'In-Person Course', 'yn', 'Classroom option at a physical location.', ('feat', 'in_person_course_available')),
    ('Delivery / Format', 'Self-Paced', 'yn', 'Students work on their own schedule (no set class times).', ('feat', 'self_paced')),
    ('Delivery / Format', 'Live Instruction', 'yn', 'Live classes or webinars with an instructor (online or in person).', ('feat', 'live_instruction_available')),
    ('Course Materials', 'Physical Book', 'yn', 'A printed textbook is included or can be ordered.', ('feat', 'physical_book_available')),
    ('Course Materials', 'PDF', 'yn', 'Course text available as a PDF / digital book.', ('feat', 'digital_book_available')),
    ('Course Materials', 'Downloadable Materials', 'yn', 'Materials (PDFs, worksheets, slides) can be downloaded and kept.', ('feat', 'downloadable_materials')),
    ('Course Materials', 'Online Videos', 'yn', 'Recorded video lessons are part of the course.', ('feat', 'video_lessons_available')),
    ('Learning Format', 'Text-Based Course', 'yn', 'Lessons are mainly reading (online text).', ('feat', 'text_based_course_available')),
    ('Learning Format', 'Audio Content', 'yn', 'Audio lessons or a listen-along option.', ('feat', 'audio_content_available')),
    ('Support', 'Phone Support', 'yn', 'A phone number students can call for help.', ('feat', 'phone_support')),
    ('Support', 'Email Support', 'yn', 'A support email address or contact form.', ('feat', 'email_support')),
    ('Support', 'Live Chat Support', 'yn', 'Live chat on the website (staffed or automated).', ('feat', 'live_chat_support')),
]
GROUP_FILL = {'School': NAVY, 'Offers & Pricing': '2F4652', 'Delivery / Format': '7E9330', 'Course Materials': '5E6F22', 'Learning Format': '4A6B7A', 'Support': '3B5664', 'Reviewer': GREY}
FIXED = [('School', '#', 5), ('School', 'School', 38), ('School', 'Website', 34), ('School', 'CTEC provider #', 13), ('School', '20-Hour course page', 30), ('School', '60-Hour course page', 30)]
TAIL = [('Reviewer', 'Notes', 40), ('Reviewer', 'Checked by', 14), ('Reviewer', 'Date checked', 13), ('Reviewer', 'Progress', 10)]

wb = Workbook()
# ---------------- Instructions ----------------
ins = wb.active; ins.title = 'Instructions'
ins.sheet_view.showGridLines = False
ins.column_dimensions['A'].width = 3; ins.column_dimensions['B'].width = 34; ins.column_dimensions['C'].width = 90
r = 2
ins.cell(r, 2, 'CTEC Provider Feature Check').font = Font(name=F, size=18, bold=True, color=NAVY); r += 1
ins.cell(r, 2, f'{len(P)} schools listed on taxcoursefinder.com. Visit each school\'s website and record what it actually offers.').font = Font(name=F, size=11, color=GREY); r += 2
steps = [
    ('1. Open the "Checklist" tab', 'One row per school. Click the website link (and the course page links where given).'),
    ('2. Fill the shaded cells only', 'Light-yellow cells are for you. Grey cells (school, website, CTEC #) are already filled in — don\'t change them.'),
    ('3. Yes / No / Unsure', 'Pick from the dropdown. Yes = the site clearly shows it. No = you looked and it isn\'t offered. Unsure = couldn\'t tell (add a note).'),
    ('4. Prices and coupons', 'Type the complete course bundle price as a number (e.g. 69.95). Leave blank if the site doesn\'t publish one. Note any coupon code or sale and its end date.'),
    ('5. Sign off each row', 'Add your initials and the date. Progress shows 100% when every Yes/No question in the row is answered.'),
    ('Reference', 'The "Currently on site" tab shows what the directory displays today, so you can see what your check will change. Check the school\'s site first; don\'t copy from it.'),
]
for a, b in steps:
    ins.cell(r, 2, a).font = Font(name=F, bold=True, color=NAVY)
    c = ins.cell(r, 3, b); c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical='top')
    ins.cell(r, 2).alignment = Alignment(vertical='top'); r += 1
r += 1
ins.cell(r, 2, 'What each column means').font = Font(name=F, size=13, bold=True, color=NAVY); r += 1
for g, h, k, d, _ in COLS:
    ins.cell(r, 2, f'{h}').font = Font(name=F, bold=True)
    c = ins.cell(r, 3, f'{g} — {d}'); c.font = Font(name=F); c.alignment = Alignment(wrap_text=True)
    r += 1
r += 1
ins.cell(r, 2, 'Example row (how a finished row looks)').font = Font(name=F, size=13, bold=True, color=NAVY); r += 1
example = {'School': 'Example Tax School (not a real listing)', 'Website': 'https://www.example.com', 'Offers 20-Hour CE': 'Yes', 'Offers 60-Hour QE': 'No',
           '20-Hour course bundle price ($)': 69.95, 'Coupon or sale (code / details)': 'RENEW10 – 10% off, ends 10/31/2026', 'Online Course': 'Yes', 'In-Person Course': 'No',
           'Self-Paced': 'Yes', 'Live Instruction': 'No', 'Physical Book': 'Yes', 'PDF': 'Yes', 'Downloadable Materials': 'Unsure', 'Online Videos': 'Yes', 'Text-Based Course': 'Yes',
           'Audio Content': 'No', 'Phone Support': 'Yes', 'Email Support': 'Yes', 'Live Chat Support': 'No', 'Notes': 'Downloads only mentioned in FAQ, not on course page', 'Checked by': 'JD', 'Date checked': '10/14/2026'}
for k, v in example.items():
    ins.cell(r, 2, k).font = Font(name=F, color=GREY)
    c = ins.cell(r, 3, v); c.font = Font(name=F, italic=True); c.alignment = Alignment(horizontal='left')
    if isinstance(v, float): c.number_format = '$#,##0.00'
    r += 1
r += 1
ins.cell(r, 2, 'Status').font = Font(name=F, size=13, bold=True, color=NAVY); r += 1
N = len(P); first, last = 3, 2 + N
prog_col = L(len(FIXED) + len(COLS) + len(TAIL))
ins.cell(r, 2, 'Schools to check').font = Font(name=F); ins.cell(r, 3, f"=COUNTA(Checklist!B{first}:B{last})").font = Font(name=F); ins.cell(r, 3).alignment = Alignment(horizontal='left'); r += 1
ins.cell(r, 2, 'Fully checked').font = Font(name=F); ins.cell(r, 3, f'=COUNTIF(Checklist!{prog_col}{first}:{prog_col}{last},1)').font = Font(name=F); ins.cell(r, 3).alignment = Alignment(horizontal='left'); r += 1
ins.cell(r, 2, 'Started').font = Font(name=F); ins.cell(r, 3, f'=COUNTIF(Checklist!{prog_col}{first}:{prog_col}{last},">0")').font = Font(name=F); ins.cell(r, 3).alignment = Alignment(horizontal='left'); r += 1
ins.cell(r + 1, 2, 'Source: school list exported from taxcoursefinder.com on Oct 12, 2026.').font = Font(name=F, size=9, color=GREY)

# ---------------- Checklist ----------------
def build(ws, prefill):
    allc = FIXED + [(g, h, 16 if k != 'text' else 30) for g, h, k, _, _ in COLS] + TAIL
    if prefill: allc = allc[:-len(TAIL)]  # no reviewer columns on the reference tab
    for i, (g, h, w) in enumerate(allc, 1):
        ws.column_dimensions[L(i)].width = w
        c = ws.cell(2, i, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=GROUP_FILL[g])
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c.border = box
    # group band row
    i = 1
    while i <= len(allc):
        g = allc[i - 1][0]; j = i
        while j < len(allc) and allc[j][0] == g: j += 1
        if j > i: ws.merge_cells(start_row=1, start_column=i, end_row=1, end_column=j)
        c = ws.cell(1, i, g.upper()); c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = PatternFill('solid', fgColor=GROUP_FILL[g]); c.alignment = Alignment(horizontal='center')
        i = j + 1
    ws.row_dimensions[2].height = 44
    for n, p in enumerate(P):
        r = 3 + n
        vals = [n + 1, p['name'], p['website'] or '', p['ctec_number'] or '', (p['ce'] or {}).get('course_url', '') if p['offers_20_hour_ce'] else '', (p['qe'] or {}).get('course_url', '') if p['offers_60_hour_qe'] else '']
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.font = Font(name=F, size=10); c.border = box; c.fill = PatternFill('solid', fgColor='F4F6F7')
            if i in (3, 5, 6) and v: c.hyperlink = v; c.font = Font(name=F, size=10, color='1F5FA8', underline='single')
        for k, (g, h, kind, _, src) in enumerate(COLS):
            col = len(FIXED) + 1 + k
            c = ws.cell(r, col); c.border = box; c.font = Font(name=F, size=10); c.alignment = Alignment(horizontal='center' if kind == 'yn' else 'left')
            if kind == 'money': c.number_format = '$#,##0.00'
            if prefill:
                t, key = src
                if t == 'flag': c.value = 'Yes' if p[key] else 'No'
                elif t == 'feat': c.value = 'Yes' if p['features'].get(key) else 'Not set'
                elif t == 'price':
                    pr = p[key] or {}
                    c.value = pr.get('current_price') if pr.get('current_price') is not None else (pr.get('display_price') or '')
                elif t == 'offers':
                    c.value = '; '.join(filter(None, [(o.get('coupon_code') or '') + (' – ' if o.get('coupon_code') else '') + (o.get('offer_title') or '') for o in p['offers']])) or ''
            else:
                c.fill = PatternFill('solid', fgColor='FFF9DB')
        if not prefill:
            base = len(FIXED) + len(COLS)
            for t in range(3):
                c = ws.cell(r, base + 1 + t); c.border = box; c.font = Font(name=F, size=10); c.fill = PatternFill('solid', fgColor='FFF9DB')
                if t == 2: c.number_format = 'mm/dd/yyyy'
            yn = [len(FIXED) + 1 + k for k, x in enumerate(COLS) if x[2] == 'yn']
            rng = ','.join(f'{L(x)}{r}' for x in yn)
            c = ws.cell(r, base + 4, f'=({"+".join(f"COUNTIF({L(x)}{r},\"?*\")" for x in yn)})/{len(yn)}')
            c.number_format = '0%'; c.font = Font(name=F, size=10, bold=True); c.border = box; c.alignment = Alignment(horizontal='center')
    last = 2 + len(P)
    ws.freeze_panes = 'C3'
    ws.auto_filter.ref = f'A2:{L(len(allc))}{last}'
    if not prefill:
        dv = DataValidation(type='list', formula1='"Yes,No,Unsure"', allow_blank=True, showErrorMessage=True, errorTitle='Pick from the list', error='Choose Yes, No or Unsure.')
        ws.add_data_validation(dv)
        for k, x in enumerate(COLS):
            col = L(len(FIXED) + 1 + k)
            if x[2] == 'yn': dv.add(f'{col}3:{col}{last}')
            if x[2] == 'money':
                dm = DataValidation(type='decimal', operator='between', formula1='0', formula2='100000', allow_blank=True, showErrorMessage=True, error='Enter the price as a number, e.g. 69.95')
                ws.add_data_validation(dm); dm.add(f'{col}3:{col}{last}')
        dd = DataValidation(type='date', operator='greaterThan', formula1='45000', allow_blank=True, showErrorMessage=True, error='Enter a date, e.g. 10/14/2026')
        ws.add_data_validation(dd); dd.add(f'{L(len(FIXED)+len(COLS)+3)}3:{L(len(FIXED)+len(COLS)+3)}{last}')
        area = f'{L(len(FIXED)+1)}3:{L(len(FIXED)+len(COLS))}{last}'
        ws.conditional_formatting.add(area, CellIsRule(operator='equal', formula=['"Yes"'], fill=PatternFill('solid', fgColor='DDEBC0'), font=Font(name=F, color='3F5A12')))
        ws.conditional_formatting.add(area, CellIsRule(operator='equal', formula=['"No"'], fill=PatternFill('solid', fgColor='F6D6D6'), font=Font(name=F, color='8A1F1F')))
        ws.conditional_formatting.add(area, CellIsRule(operator='equal', formula=['"Unsure"'], fill=PatternFill('solid', fgColor='FBE3B0'), font=Font(name=F, color='7A5200')))
        pc = L(len(allc))
        ws.conditional_formatting.add(f'{pc}3:{pc}{last}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='DDEBC0')))
        for k, x in enumerate(COLS):
            ws.cell(2, len(FIXED) + 1 + k).comment = Comment(x[3], 'Tax Course Finder')
    ws.sheet_properties.tabColor = OLIVE if not prefill else 'B0BEC5'

build(wb.create_sheet('Checklist'), False)
ref = wb.create_sheet('Currently on site')
build(ref, True)
ins.sheet_properties.tabColor = NAVY
wb.active = 1
out = S + '/xlsx/CTEC-Provider-Feature-Checklist.xlsx'
wb.save(out); print(out)
